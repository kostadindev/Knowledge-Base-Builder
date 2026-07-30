import logging
from knowledge_base_builder.base_processor import BaseProcessor

logger = logging.getLogger(__name__)

try:
    from markitdown import MarkItDown
    _MARKITDOWN_AVAILABLE = True
except ImportError:
    _MARKITDOWN_AVAILABLE = False

class PresentationProcessor(BaseProcessor):
    SUPPORTED_EXTENSIONS = ['.pptx']

    @staticmethod
    def download(url: str) -> str:
        return BaseProcessor.download(url, PresentationProcessor.SUPPORTED_EXTENSIONS)

    @staticmethod
    def extract_text(file_path: str) -> str:
        # Try MarkItDown first if available
        if _MARKITDOWN_AVAILABLE:
            try:
                md = MarkItDown()
                result = md.convert(file_path)
                logger.info("Extracted PPTX with MarkItDown")
                return result.text_content
            except Exception as e:
                logger.warning("MarkItDown failed for PPTX, falling back: %s", e)

        # Fallback to python-pptx
        from pptx import Presentation
        logger.info("Extracting PPTX with python-pptx")
        prs = Presentation(file_path)

        parts = []
        for i, slide in enumerate(prs.slides, 1):
            slide_text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_text.append(text)

            # Extract speaker notes
            notes = ''
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()

            if slide_text or notes:
                parts.append(f"## Slide {i}")
                if slide_text:
                    parts.append('\n'.join(slide_text))
                if notes:
                    parts.append(f"\n*Speaker Notes:* {notes}")

        return '\n\n'.join(parts)
