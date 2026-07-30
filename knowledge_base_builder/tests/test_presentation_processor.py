"""Tests for PresentationProcessor — real .pptx creation + MarkItDown fallback."""

import os
import sys
import tempfile
import pytest
from unittest.mock import patch, MagicMock
from knowledge_base_builder.presentation_processor import PresentationProcessor


# ---------------------------------------------------------------------------
# Real .pptx file tests (using python-pptx to create test files)
# ---------------------------------------------------------------------------
class TestRealPptx:
    @patch("knowledge_base_builder.presentation_processor._MARKITDOWN_AVAILABLE", False)
    def test_extract_from_real_pptx(self):
        """Create a real .pptx file and verify extraction."""
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        # Slide 1: title + subtitle
        layout = prs.slide_layouts[0]  # Title slide
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = "Knowledge Base Builder"
        slide.placeholders[1].text = "A Python Package for LLM-Powered KB Construction"

        # Slide 2: content slide
        layout2 = prs.slide_layouts[1]  # Title and content
        slide2 = prs.slides.add_slide(layout2)
        slide2.shapes.title.text = "Features"
        slide2.placeholders[1].text = "Multi-source ingestion\nLLM summarization\nIncremental builds"

        with tempfile.NamedTemporaryFile(suffix='.pptx', delete=False) as f:
            prs.save(f.name)
            path = f.name

        try:
            result = PresentationProcessor.extract_text(path)
            assert "Knowledge Base Builder" in result
            assert "Python Package" in result
            assert "Features" in result
            assert "Multi-source ingestion" in result
            assert "## Slide 1" in result
            assert "## Slide 2" in result
        finally:
            os.unlink(path)

    @patch("knowledge_base_builder.presentation_processor._MARKITDOWN_AVAILABLE", False)
    def test_extract_with_speaker_notes(self):
        """Test extraction of speaker notes from real .pptx."""
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Slide With Notes"
        # Add speaker notes
        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = "Remember to mention the API key setup"

        with tempfile.NamedTemporaryFile(suffix='.pptx', delete=False) as f:
            prs.save(f.name)
            path = f.name

        try:
            result = PresentationProcessor.extract_text(path)
            assert "Slide With Notes" in result
            assert "Speaker Notes:" in result
            assert "API key setup" in result
        finally:
            os.unlink(path)

    @patch("knowledge_base_builder.presentation_processor._MARKITDOWN_AVAILABLE", False)
    def test_empty_presentation(self):
        """Presentation with no slides."""
        from pptx import Presentation
        prs = Presentation()
        with tempfile.NamedTemporaryFile(suffix='.pptx', delete=False) as f:
            prs.save(f.name)
            path = f.name
        try:
            result = PresentationProcessor.extract_text(path)
            assert result == ""
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# MarkItDown integration
# ---------------------------------------------------------------------------
class TestMarkItDownIntegration:
    @patch("knowledge_base_builder.presentation_processor._MARKITDOWN_AVAILABLE", True)
    @patch("knowledge_base_builder.presentation_processor.MarkItDown", create=True)
    def test_markitdown_used_when_available(self, MockMD):
        mock_result = MagicMock()
        mock_result.text_content = "# Slide 1\nConverted content"
        MockMD.return_value.convert.return_value = mock_result

        result = PresentationProcessor.extract_text("test.pptx")
        assert result == "# Slide 1\nConverted content"
        MockMD.return_value.convert.assert_called_once_with("test.pptx")

    @patch("knowledge_base_builder.presentation_processor._MARKITDOWN_AVAILABLE", True)
    @patch("knowledge_base_builder.presentation_processor.MarkItDown", create=True)
    def test_markitdown_error_falls_back_to_pptx(self, MockMD):
        """When MarkItDown raises, fall back to python-pptx."""
        MockMD.return_value.convert.side_effect = Exception("MarkItDown crash")

        from pptx import Presentation
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Fallback Content"
        with tempfile.NamedTemporaryFile(suffix='.pptx', delete=False) as f:
            prs.save(f.name)
            path = f.name
        try:
            result = PresentationProcessor.extract_text(path)
            assert "Fallback Content" in result
        finally:
            os.unlink(path)


# ---------------------------------------------------------------------------
# Download delegation
# ---------------------------------------------------------------------------
class TestDownload:
    @patch("knowledge_base_builder.base_processor.BaseProcessor.download")
    def test_download_delegates_to_base(self, mock_dl):
        mock_dl.return_value = "/tmp/slides.pptx"
        result = PresentationProcessor.download("https://example.com/slides.pptx")
        mock_dl.assert_called_once_with("https://example.com/slides.pptx", ['.pptx'])
        assert result == "/tmp/slides.pptx"
